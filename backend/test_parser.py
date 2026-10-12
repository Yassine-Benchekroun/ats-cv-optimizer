import json
from document_parser import extract_text_from_file
import docx
from fpdf import FPDF
import os

def create_sample_docx(file_path):
    doc = docx.Document()
    doc.add_heading('John Doe - Resume', 0)
    doc.add_paragraph('Highly motivated software engineer with 5 years of experience in Python, FastAPI, and data engineering.')
    doc.add_heading('Experience', level=1)
    doc.add_paragraph('Software Engineer at Tech Innovations Inc.\nDeveloped scalable microservices.')
    doc.add_heading('Education', level=1)
    doc.add_paragraph('B.S. in Computer Science - University of State')
    doc.save(file_path)

def create_sample_pdf(file_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=16)
    pdf.cell(200, 10, text="John Doe - Resume", align='C')
    pdf.ln(20)
    pdf.set_font("helvetica", size=12)
    pdf.multi_cell(0, 10, text="Highly motivated software engineer with 5 years of experience in Python, FastAPI, and data engineering.")
    pdf.ln(10)
    pdf.set_font("helvetica", size=14)
    pdf.cell(200, 10, text="Experience")
    pdf.ln(10)
    pdf.set_font("helvetica", size=12)
    pdf.multi_cell(0, 10, text="Software Engineer at Tech Innovations Inc.\nDeveloped scalable microservices.")
    pdf.ln(10)
    pdf.set_font("helvetica", size=14)
    pdf.cell(200, 10, text="Education")
    pdf.ln(10)
    pdf.set_font("helvetica", size=12)
    pdf.multi_cell(0, 10, text="B.S. in Computer Science - University of State")
    pdf.output(file_path)

if __name__ == "__main__":
    docx_path = "test_resume.docx"
    pdf_path = "test_resume.pdf"
    
    print("Generating mock DOCX and PDF files...")
    create_sample_docx(docx_path)
    create_sample_pdf(pdf_path)
    
    print("\n--- Testing DOCX Extraction ---")
    docx_result = extract_text_from_file(docx_path)
    print(json.dumps(docx_result, indent=2))
    
    print("\n--- Testing PDF Extraction ---")
    pdf_result = extract_text_from_file(pdf_path)
    print(json.dumps(pdf_result, indent=2))
    
    print("\n--- Testing Invalid File ---")
    invalid_result = extract_text_from_file("nonexistent.txt")
    print(json.dumps(invalid_result, indent=2))
    
    # Cleanup (optional)
    if os.path.exists(docx_path): os.remove(docx_path)
    if os.path.exists(pdf_path): os.remove(pdf_path)
